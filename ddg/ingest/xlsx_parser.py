"""XLSX parsing into anchored cell nodes.

openpyxl stores formulas but does not evaluate them [ref 47]. The workbook is
therefore read twice - once for formulas, once for Excel's last cached results -
and every cached value is flagged. A cached result is never presented as a
freshly recalculated one.
"""

from __future__ import annotations

from pathlib import Path

import openpyxl

from ddg.anchors import make_quote
from ddg.models import LocationSelector, Node, NodeKind, QuantityContext, SourceSnapshot
from ddg.units import parse_decimal

#: Functions the deterministic checkers do not implement. Their presence is
#: reported so affected checks abstain rather than guess.
UNSUPPORTED_FUNCTIONS = ("INDIRECT", "OFFSET", "RAND", "NOW", "TODAY", "WEBSERVICE")


def parse_xlsx(path: Path, snap: SourceSnapshot) -> tuple[SourceSnapshot, list[Node]]:
    wb_f = openpyxl.load_workbook(str(path), data_only=False, read_only=False)
    wb_v = openpyxl.load_workbook(str(path), data_only=True, read_only=False)

    unsupported: list[str] = []
    if path.suffix.lower() == ".xlsm" or getattr(wb_f, "vba_archive", None):
        unsupported.append("workbook contains macros; macros are never executed (N01)")
    ext = getattr(wb_f, "_external_links", None)
    if ext:
        unsupported.append(
            f"{len(ext)} external workbook link(s); values may be stale and are not refreshed"
        )

    nodes: list[Node] = []
    for ws_f in wb_f.worksheets:
        ws_v = wb_v[ws_f.title]
        for row in ws_f.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                addr = cell.coordinate
                formula = None
                raw = cell.value
                cached = ws_v[addr].value

                if isinstance(raw, str) and raw.startswith("="):
                    formula = raw
                    for fn in UNSUPPORTED_FUNCTIONS:
                        if fn in raw.upper():
                            unsupported.append(
                                f"{ws_f.title}!{addr} uses {fn}, which the checkers do not evaluate"
                            )
                    value_source, is_cached = cached, True
                else:
                    value_source, is_cached = raw, False

                text = str(value_source) if value_source is not None else str(raw)
                normalized = (
                    parse_decimal(text) if not isinstance(value_source, bool) else None
                )

                selector = LocationSelector(
                    document_id=snap.document_id,
                    kind=NodeKind.SHEET_CELL,
                    quote=make_quote(text, 0, len(text)) if text else None,
                    sheet_name=ws_f.title,
                    cell_ref=addr,
                    structural_path=(ws_f.title,),
                )
                nodes.append(Node(
                    node_id=f"{snap.document_id}#{ws_f.title}!{addr}",
                    source_version=snap.version_hash,
                    kind=NodeKind.SHEET_CELL,
                    selector=selector,
                    evidence_text=text or "(empty)",
                    raw_value=str(raw),
                    normalized_value=normalized,
                    is_cached_value=is_cached,
                    formula=formula,
                    context=QuantityContext(),
                ))

    wb_f.close()
    wb_v.close()
    return snap.model_copy(update={"unsupported_objects": unsupported}), nodes
