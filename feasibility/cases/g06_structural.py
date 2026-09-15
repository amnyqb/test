"""G06 - Do structural edits keep every link on the right content?

One micro-test is a small report and workbook describing the same line items,
reviewed so that each sentence is linked to its workbook cell. One to three
structural edits are then applied, and the revision goes through `ddg revise`:

* workbook: a row inserted, deleted, moved, relabelled, or given another row's
  label;
* report: a sentence moved, deleted or repeated elsewhere; a section inserted,
  which renumbers every section after it.

Every line item carries an identity the documents never show. Each re-anchored
node is judged against it:

* followed      - it now points at the same item (or at an identical repeat of
  the same sentence);
* lost          - the item is still there, but the link went to review (a cost);
* removed       - the item was deleted and the link is unresolved (correct);
* flagged       - it points at different content, but its meaning is already
  flagged for review, so no check runs on it (a cost and a risk, not silent);
* indistinguishable - it moved to a row with the identical label and value, so
  nothing in either document could tell the two apart;
* MISATTACHED silently - different content, unflagged: critical.

Values never change, so every relation that still executes must PASS. Few
distinct values are used, so neighbouring rows often hold the same number -
the condition under which positional matching misattaches without any sign.
"""

from __future__ import annotations

import copy
import json
import random
import tempfile
from collections import Counter
from dataclasses import dataclass, field
from decimal import Decimal
from pathlib import Path

import docx
import openpyxl

from ddg.models import CheckStatus, NodeKind
from ddg.pipeline import audit
from ddg.revision.engine import revise
from ddg.store import Store
from feasibility.harness import FailureKind, Outcome, UseCase

MARK = "SYNTHETIC FIXTURE - NOT REAL PROJECT MATERIAL"
LABELS = [
    "Civil works", "Equipment", "Contingency", "Design fees", "Grid connection", "Land",
    "Insurance", "Permits", "Commissioning", "Spares", "Training", "Legal", "Financing fees",
    "Owner costs", "Site security", "Water supply", "Access roads", "Transformers", "Cabling",
    "Control systems",
]
SECTION_TITLES = ["Capital costs", "Infrastructure", "Owner costs", "Allowances", "Services",
                  "Delivery", "Support"]
MILLIONS = [Decimal(v) for v in ("1.2", "2.5", "3.1", "4.0", "6.2", "8.0")]


@dataclass
class Item:
    uid: int
    label: str      # the workbook row label, which edits may change
    name: str       # how the report refers to the item, which never changes
    millions: Decimal


@dataclass
class Section:
    uid: int
    title: str
    mentions: list[tuple[int, int]] = field(default_factory=list)  # (item uid, repeat no.)
    filler: bool = False


@dataclass
class Package:
    rows: list[Item]
    sections: list[Section]
    catalog: dict[int, Item]
    next_uid: int


def _package(rng: random.Random) -> Package:
    rows: list[Item] = []
    for uid, label in enumerate(rng.sample(LABELS, rng.randint(3, 7))):
        rows.append(Item(uid, label, label, rng.choice(MILLIONS)))
    order = [item.uid for item in rows]
    rng.shuffle(order)
    sections: list[Section] = []
    next_uid = len(rows)
    while order:
        take = rng.randint(1, 3)
        sections.append(Section(next_uid, rng.choice(SECTION_TITLES),
                                [(uid, 0) for uid in order[:take]]))
        next_uid += 1
        order = order[take:]
    return Package(rows, sections, {item.uid: item for item in rows}, next_uid)


def _render(pkg: Package, root: Path) -> dict[str, tuple]:
    """Write model.xlsx and report.docx; return node id -> hidden identity."""
    root.mkdir(parents=True)
    identity: dict[str, tuple] = {}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"] = MARK
    first = 3
    for r, item in enumerate(pkg.rows, start=first):
        ws[f"A{r}"] = item.label
        ws[f"B{r}"] = int(item.millions * 1_000_000)
        identity[f"model.xlsx#Model!B{r}"] = ("cell", item.uid)
    last = first + len(pkg.rows) - 1
    ws[f"A{last + 2}"] = "Total"
    ws[f"B{last + 2}"] = f"=SUM(B{first}:B{last})"
    identity[f"model.xlsx#Model!B{last + 2}"] = ("total",)
    wb.save(str(root / "model.xlsx"))

    d = docx.Document()
    d.add_heading("Feasibility Report (SYNTHETIC)", level=1)
    d.add_paragraph(MARK)
    index = 2
    for number, section in enumerate(pkg.sections, start=1):
        d.add_heading(f"{number}. {section.title}", level=2)
        index += 1
        if section.filler:
            d.add_paragraph("Details for this section follow in a later revision.")
            index += 1
        for uid, repeat in section.mentions:
            item = pkg.catalog[uid]
            d.add_paragraph(f"{item.name} is budgeted at USD {item.millions} million in FY26.")
            identity[f"report.docx#p{index}q0"] = ("mention", uid, repeat)
            index += 1
    d.save(str(root / "report.docx"))
    return identity


def _manifest(identity: dict[str, tuple]) -> dict:
    base = dict(entity="ProjectCo", period="FY26", scenario="base", unit="currency",
                currency="USD")
    cell_of = {ident[1]: nid for nid, ident in identity.items() if ident[0] == "cell"}
    contexts: dict[str, dict] = {}
    relations: list[dict] = []
    for nid, ident in identity.items():
        if ident[0] == "cell":
            contexts[nid] = {**base, "metric": f"item{ident[1]}", "scale": "units"}
        elif ident[0] == "mention":
            contexts[nid] = {**base, "metric": f"item{ident[1]}", "scale": "millions"}
            relations.append({
                "relation_id": f"rel:item{ident[1]}:repeat{ident[2]}",
                "type": "SAME_QUANTITY_AS",
                "endpoints": [{"role": "narrative", "node_id": nid},
                              {"role": "workbook", "node_id": cell_of[ident[1]]}],
                "reviewer": "generator", "reason": "linked by construction",
            })
    return {"contexts": contexts, "relations": relations}


def _mutate(rng: random.Random, pkg: Package) -> list[str]:
    edits: list[str] = []
    for _ in range(rng.randint(1, 3)):
        kind = rng.choice(["insert_row", "delete_row", "move_row", "rename_label",
                           "duplicate_label", "insert_section", "move_mention",
                           "delete_mention", "repeat_mention"])
        rows = pkg.rows
        unused = [l for l in LABELS if l not in {i.label for i in pkg.catalog.values()}]
        mentioned = [s for s in pkg.sections if s.mentions]

        if kind == "insert_row" and unused:
            label = rng.choice(unused)
            item = Item(pkg.next_uid, label, label, rng.choice(MILLIONS))
            pkg.next_uid += 1
            pkg.catalog[item.uid] = item
            position = rng.randint(0, len(rows))
            rows.insert(position, item)
            edits.append(f"insert row {label!r} ({item.millions}m) at {position}")
        elif kind == "delete_row" and len(rows) > 2:
            item = rows.pop(rng.randrange(len(rows)))
            edits.append(f"delete row {item.label!r}")
        elif kind == "move_row" and len(rows) > 1:
            item = rows.pop(rng.randrange(len(rows)))
            position = rng.randint(0, len(rows))
            rows.insert(position, item)
            edits.append(f"move row {item.label!r} to {position}")
        elif kind == "rename_label" and unused:
            item = rng.choice(rows)
            old, item.label = item.label, rng.choice(unused)
            edits.append(f"relabel row {old!r} as {item.label!r}")
        elif kind == "duplicate_label" and len(rows) > 1:
            kept, changed = rng.sample(rows, 2)
            edits.append(f"relabel row {changed.label!r} as duplicate {kept.label!r}")
            changed.label = kept.label
        elif kind == "insert_section":
            position = rng.randint(0, len(pkg.sections))
            pkg.sections.insert(position, Section(pkg.next_uid, rng.choice(SECTION_TITLES),
                                                  filler=True))
            pkg.next_uid += 1
            edits.append(f"insert section at {position}, renumbering those after it")
        elif kind == "move_mention" and mentioned:
            source = rng.choice(mentioned)
            mention = source.mentions.pop(rng.randrange(len(source.mentions)))
            target = rng.choice(pkg.sections)
            target.mentions.insert(rng.randint(0, len(target.mentions)), mention)
            edits.append(f"move sentence for item {mention[0]} (repeat {mention[1]})")
        elif kind == "delete_mention" and mentioned:
            source = rng.choice(mentioned)
            mention = source.mentions.pop(rng.randrange(len(source.mentions)))
            edits.append(f"delete sentence for item {mention[0]} (repeat {mention[1]})")
        elif kind == "repeat_mention" and mentioned:
            uid, _ = rng.choice(rng.choice(mentioned).mentions)
            repeat = 1 + max(r for s in pkg.sections for u, r in s.mentions if u == uid)
            target = rng.choice(pkg.sections)
            target.mentions.insert(rng.randint(0, len(target.mentions)), (uid, repeat))
            edits.append(f"repeat the sentence for item {uid} in another place")
    return edits


def _indistinguishable(change) -> bool:
    """Identical label and value: no evidence in either document separates the rows."""
    return (change.old.kind is NodeKind.SHEET_CELL and change.new.kind is NodeKind.SHEET_CELL
            and change.old.selector.structural_path == change.new.selector.structural_path
            and change.old.raw_value == change.new.raw_value)


def run_one(rng: random.Random) -> Outcome:
    before_pkg = _package(rng)
    after_pkg = copy.deepcopy(before_pkg)
    edits = _mutate(rng, after_pkg)

    with tempfile.TemporaryDirectory(prefix="g06-") as tmp:
        root = Path(tmp)
        before = _render(before_pkg, root / "v1")
        (root / "v1" / "review.json").write_text(json.dumps(_manifest(before)))
        after = _render(after_pkg, root / "v2")
        with Store() as store:
            baseline = audit(root / "v1", store)
            rep = revise(root / "v2", store)

    if baseline.review_problems:
        raise RuntimeError(f"generator produced an invalid manifest: {baseline.review_problems}")
    unclean = [r.check_id for r in baseline.results
               if r.check_id.startswith("chk:rel:item") and r.status is not CheckStatus.PASS]
    if unclean:
        raise RuntimeError(f"baseline audit is not clean: {unclean}")

    where_now: dict[tuple, list[str]] = {}
    for nid, ident in after.items():
        where_now.setdefault(ident, []).append(nid)

    counts: Counter[str] = Counter()
    misattached: list[str] = []
    for change in rep.changes:
        if change.old is None or change.old.node_id not in before:
            continue
        ident = before[change.old.node_id]
        if change.mapped:
            landed = after.get(change.new.node_id)
            if landed == ident:
                counts["links followed"] += 1
            elif landed is not None and landed[0] == "mention" and landed[:2] == ident[:2]:
                counts["links followed to an identical repeat"] += 1
            elif _indistinguishable(change):
                counts["links moved to a row the documents cannot tell apart"] += 1
            elif change.context_suspect:
                counts["links misattached but flagged for review"] += 1
            else:
                counts["links MISATTACHED silently"] += 1
                misattached.append(f"{change.old.node_id} {ident} -> {change.new.node_id} {landed}")
        elif ident in where_now:
            counts["links lost to review"] += 1
        else:
            counts["links unresolved, content removed"] += 1
    for o in rep.closure.outcomes:
        counts[f"relations {o.disposition.value}"] += 1

    unexpected = [f"{r.check_id} {r.status.value}" for r in rep.results
                  if r.check_id.startswith("chk:rel:item") and r.status is not CheckStatus.PASS]
    size = len(edits) + len(before_pkg.rows)
    detail = (f"edits: {'; '.join(edits) or 'none applicable'} | "
              f"misattached: {'; '.join(misattached[:3]) or '-'} | "
              f"unexpected verdicts: {'; '.join(unexpected[:3]) or '-'} | "
              f"closure misses: {len(rep.closure_misses)}")

    if misattached:
        return Outcome(False, "silent misattachment", critical=True, size=size,
                       detail=detail, counts=dict(counts))
    if rep.closure_misses:
        return Outcome(False, "closure miss", size=size, detail=detail, counts=dict(counts))
    if unexpected:
        return Outcome(False, "unexpected verdict on unchanged values", size=size,
                       detail=detail, counts=dict(counts))
    label = ("no misattachment; nothing sent to review" if not counts["links lost to review"]
             else "no misattachment; some links sent to review")
    return Outcome(True, label, size=size, detail=detail, counts=dict(counts))


USE_CASE = UseCase(
    id="G06",
    question="Do structural edits keep every link on the right content?",
    failure_kind=FailureKind.DESIGN,
    critical_means="any link silently re-attached to different content",
    run_one=run_one,
    default_n=400,
)
