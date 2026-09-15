"""Regression tests for defects confirmed by feasibility probes on 15 September 2026.

Each test pins a defect a quick probe found before any test harness existed:
identity judged without metric, a missing scale assumed to be units, no
rounding, sums that ignored context, misread prose figures, and report table
cells without headers. The two review loopholes are pinned in
test_review_queue.py. All inputs are synthetic.
"""

from __future__ import annotations

from decimal import Decimal

import docx
import pytest

from ddg import CHECKER_VERSION
from ddg.check import run_check
from ddg.ingest import ingest_file
from ddg.models import (
    CheckDefinition,
    CheckResult,
    CheckStatus,
    Endpoint,
    ExtractionMethod,
    LocationSelector,
    Node,
    NodeKind,
    QuantityContext,
    Relation,
    RelationType,
    ReviewState,
)

CTX = dict(entity="ProjectCo", metric="capex", period="FY26", scenario="base",
           unit="currency", currency="USD")


def node(node_id: str, value: str, *, doc: str = "d", **context) -> Node:
    return Node(
        node_id=node_id, source_version="v1", kind=NodeKind.PARAGRAPH,
        selector=LocationSelector(document_id=doc, kind=NodeKind.PARAGRAPH),
        evidence_text="synthetic", normalized_value=Decimal(value),
        context=QuantityContext(**{**CTX, **context}),
    )


def run(rtype: RelationType, participants: list[tuple[str, Node]]) -> CheckResult:
    nodes = {n.node_id: n for _, n in participants}
    relation = Relation(
        relation_id="r", type=rtype,
        endpoints=tuple(Endpoint(role=role, node_id=n.node_id) for role, n in participants),
        extraction_method=ExtractionMethod.HUMAN_AUTHORED,
        review_state=ReviewState.ACCEPTED, valid_source_versions=("v1",),
    )
    check = CheckDefinition(check_id="c", relation_id="r", expression=rtype.value,
                            operand_node_ids=tuple(nodes), tolerance=Decimal("0.005"),
                            checker_version=CHECKER_VERSION)
    return run_check(check, relation, nodes, "run")


def pair(a: Node, b: Node) -> CheckResult:
    return run(RelationType.SAME_QUANTITY_AS, [("a", a), ("b", b)])


# --- identity: equal values are not enough ---------------------------------

def test_equal_values_with_different_metrics_are_not_the_same_quantity():
    res = pair(node("revenue", "8.0", scale="millions", metric="revenue"),
               node("contingency", "8000000", scale="units", metric="contingency"))
    assert res.status is CheckStatus.NEEDS_REVIEW
    assert "metric differs" in res.detail


# --- scale: never assumed ---------------------------------------------------

@pytest.mark.parametrize("value, scale", [("12400000", "units"), ("12.4", "units")])
def test_a_missing_scale_is_never_assumed_to_be_units(value, scale):
    res = pair(node("narrative", "12.4", scale=None), node("cell", value, scale=scale))
    assert res.status is CheckStatus.NEEDS_REVIEW
    assert "missing scale is never assumed" in res.detail


def test_undeclared_scales_from_different_documents_need_review():
    res = pair(node("a", "12.4", scale=None, doc="report"),
               node("b", "12.4", scale=None, doc="model"))
    assert res.status is CheckStatus.NEEDS_REVIEW


def test_undeclared_scales_within_one_document_are_compared_as_written():
    res = pair(node("a", "12.4", scale=None), node("b", "12.4", scale=None))
    assert res.status is CheckStatus.PASS


# --- rounding ---------------------------------------------------------------

def test_a_rounded_figure_matches_the_exact_value_it_rounds():
    res = pair(node("narrative", "12.4", scale="millions"),
               node("cell", "12437210", scale="units"))
    assert res.status is CheckStatus.PASS
    assert any("precision of the figure" in t for t in res.computation_trace)


def test_a_difference_beyond_the_displayed_precision_still_fails():
    res = pair(node("narrative", "12.4", scale="millions"),
               node("cell", "12460000", scale="units"))
    assert res.status is CheckStatus.FAIL


def test_rounding_never_rescues_a_different_quantity():
    res = pair(node("narrative", "12.4", scale="millions"),
               node("cell", "12437210", scale="units", metric="opex"))
    assert res.status is CheckStatus.NEEDS_REVIEW


@pytest.mark.parametrize("policy, status", [
    ("exact", CheckStatus.FAIL),
    ("nearest:0.5", CheckStatus.PASS),
    ("banker's", CheckStatus.NEEDS_REVIEW),
])
def test_a_declared_rounding_policy_overrides_displayed_precision(policy, status):
    res = pair(node("narrative", "12.4", scale="millions", rounding_policy=policy),
               node("cell", "12600000", scale="units"))
    assert res.status is status


# --- sums and ratios --------------------------------------------------------

def test_a_rounded_narrative_total_matches_exact_components():
    res = run(RelationType.SUM_OF, [
        ("total", node("total", "12.4", scale="millions")),
        ("operand", node("civil", "6218605", scale="units", metric="civil")),
        ("operand", node("equipment", "6218605", scale="units", metric="equipment")),
    ])
    assert res.status is CheckStatus.PASS


def test_a_sum_whose_participants_contradict_each_other_needs_review():
    """Before the repair this passed: sums never looked at context."""
    res = run(RelationType.SUM_OF, [
        ("total", node("total", "9", period="FY26")),
        ("operand", node("a", "4", period="FY26")),
        ("operand", node("b", "5", period="FY27")),
    ])
    assert res.status is CheckStatus.NEEDS_REVIEW
    assert "period differs" in res.detail


@pytest.mark.parametrize("reported, status", [("3.1", CheckStatus.PASS),
                                              ("3.2", CheckStatus.FAIL)])
def test_a_ratio_allows_only_the_rounding_of_the_reported_figure(reported, status):
    res = run(RelationType.RATIO_OF, [
        ("result", node("gearing", reported, metric="gearing", unit="ratio", currency=None)),
        ("numerator", node("capex", "12437210")),
        ("denominator", node("equity", "4000000", metric="equity")),
    ])
    assert res.status is status


# --- reading figures in prose ------------------------------------------------

def _mentions(tmp_path, *paragraphs: str) -> list[Node]:
    d = docx.Document()
    for text in paragraphs:
        d.add_paragraph(text)
    path = tmp_path / "probe.docx"
    d.save(str(path))
    _, nodes = ingest_file(path, "pkg", tmp_path / "blobs")
    return [n for n in nodes if n.kind is NodeKind.PARAGRAPH and n.raw_value is not None]


@pytest.mark.parametrize("sentence, value, scale, currency, unit", [
    ("Capex is USD 12.4m in total.", "12.4", "millions", "USD", "currency"),
    ("Capex is $12.4m in total.", "12.4", "millions", "USD", "currency"),
    ("The fund holds £3bn of assets.", "3", "billions", "GBP", "currency"),
    ("The loss was (3.1) million this year.", "-3.1", "millions", None, None),
    ("The loss was USD (3.1) million this year.", "-3.1", "millions", "USD", "currency"),
    ("Total capex (USD 12.4 million) is fixed.", "12.4", "millions", "USD", "currency"),
    ("Net debt fell by −3.1 million.", "-3.1", "millions", None, None),
    ("The margin is 12.4% before tax.", "12.4", "percent", None, "percent"),
    ("Growth was 4 per cent last year.", "4", "percent", None, "percent"),
    ("The tower is 12.4m tall.", "12.4", None, None, None),
])
def test_prose_figures_keep_their_sign_scale_currency_and_unit(
        tmp_path, sentence, value, scale, currency, unit):
    found = _mentions(tmp_path, sentence)
    assert len(found) == 1, [n.raw_value for n in found]
    ctx = found[0].context
    assert (str(found[0].normalized_value), ctx.scale, ctx.currency, ctx.unit) == \
        (value, scale, currency, unit)


def test_look_alike_numbers_stay_out(tmp_path):
    assert _mentions(tmp_path,
                     "See section 9 and clause (3) of the FY26 plan, dated 2026.",
                     "COVID-19 delayed items 4-5.") == []


def test_a_range_dash_is_not_a_minus_sign(tmp_path):
    found = _mentions(tmp_path, "Capex of 10.5-12.5 million is expected.")
    assert sorted(str(n.normalized_value) for n in found) == ["10.5", "12.5"]


# --- report tables -----------------------------------------------------------

def test_table_cells_carry_their_headers_currency_and_scale(tmp_path):
    d = docx.Document()
    table = d.add_table(rows=3, cols=3)
    for (r, c), text in {
        (0, 1): "FY26 (USD million)", (0, 2): "FY27 (USD '000)",
        (1, 0): "Capex", (1, 1): "12.4", (1, 2): "13,100",
        (2, 0): "Margin %", (2, 1): "4.5",
    }.items():
        table.cell(r, c).text = text
    path = tmp_path / "table.docx"
    d.save(str(path))
    _, nodes = ingest_file(path, "pkg", tmp_path / "blobs")
    cells = {n.node_id.split("#")[1]: n for n in nodes if n.kind is NodeKind.TABLE_CELL}

    capex = cells["t0r1c1"]
    assert capex.selector.structural_path[-2:] == ("row=Capex", "col=FY26 (USD million)")
    assert (capex.context.currency, capex.context.scale) == ("USD", "millions")
    assert capex.context.period is None, "a header saying FY26 is a label, not an inferred period"
    assert cells["t0r1c2"].context.scale == "thousands"
    margin = cells["t0r2c1"]
    assert (margin.context.unit, margin.context.scale) == ("percent", "percent")
