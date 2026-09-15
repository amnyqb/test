"""Regressions for failures found by the G06 feasibility harness (15 Sep 2026).

Each test is the smallest form of a failure the harness generated: a figure
captured by another sentence, a sentence replaced in place, a label handed to a
different row, two rows swapping labels, and one inserted row withdrawing every
reviewed link on a sheet. The last re-anchoring test keeps the legitimate case -
a row that simply moved - followed. All inputs are synthetic.
"""

from __future__ import annotations

from pathlib import Path

import docx
import openpyxl

from ddg.ingest import ingest_file
from ddg.ingest.docx_parser import document_text
from ddg.models import (
    Endpoint,
    ExtractionMethod,
    LocationSelector,
    Node,
    NodeKind,
    Relation,
    RelationType,
    ReviewState,
)
from ddg.revision.closure import Disposition, compute_closure
from ddg.revision.remap import ChangeStatus, NodeChange, map_document


def _report(path: Path, paragraphs: list[str]) -> Path:
    d = docx.Document()
    for text in paragraphs:
        d.add_paragraph(text)
    d.save(str(path))
    return path


def _model(path: Path, rows: list[tuple[str, int]]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    for r, (label, value) in enumerate(rows, start=3):
        ws[f"A{r}"], ws[f"B{r}"] = label, value
    wb.save(str(path))
    return path


def _changes(tmp_path: Path, name: str, before, after) -> dict[str, NodeChange]:
    make = _report if name.endswith(".docx") else _model
    for version in ("v1", "v2"):
        (tmp_path / version).mkdir()
    _, old_nodes = ingest_file(make(tmp_path / "v1" / name, before), "pkg", tmp_path / "blobs")
    snap, new_nodes = ingest_file(make(tmp_path / "v2" / name, after), "pkg", tmp_path / "blobs")
    text = document_text(Path(snap.blob_path)) if name.endswith(".docx") else None
    return {c.old.node_id: c for c in map_document(old_nodes, new_nodes, text) if c.old}


# --- figures in prose -------------------------------------------------------

def test_a_repeated_sentence_nearby_never_captures_another_sentences_figure(tmp_path):
    """Seed G06:0:14: context from the next paragraph outvoted the sentence's subject."""
    changes = _changes(tmp_path, "report.docx", [
        "Owner costs is budgeted at USD 3.1 million in FY26.",
        "2. Delivery",
        "Legal is budgeted at USD 3.1 million in FY26.",
    ], [
        "Owner costs is budgeted at USD 3.1 million in FY26.",
        "Legal is budgeted at USD 3.1 million in FY26.",
        "2. Delivery",
        "Legal is budgeted at USD 3.1 million in FY26.",
    ])
    owner = changes["report.docx#p0q0"]
    assert owner.mapped and owner.new.node_id == "report.docx#p0q0"
    legal = changes["report.docx#p2q0"]
    assert not legal.mapped or legal.new.evidence_text.startswith("Legal")


def test_a_sentence_replaced_by_a_different_statement_leaves_its_figure_unresolved(tmp_path):
    """Seed G06:0:176: another sentence's figure was taken for an edited value."""
    changes = _changes(tmp_path, "report.docx", [
        "Alpha is budgeted at USD 4.0 million.",
        "Beta is budgeted at USD 2.5 million.",
        "Gamma is budgeted at USD 6.2 million.",
    ], [
        "Alpha is budgeted at USD 4.0 million.",
        "Delta is budgeted at USD 3.1 million.",
        "Gamma is budgeted at USD 6.2 million.",
    ])
    assert changes["report.docx#p1q0"].status is ChangeStatus.UNRESOLVED
    assert changes["report.docx#p0q0"].status is ChangeStatus.UNCHANGED


# --- workbook labels ------------------------------------------------------------

def test_a_label_handed_to_another_row_is_not_followed(tmp_path):
    """Seeds G06:0:396 and G06:0:191: a row was renamed and its label reused."""
    changes = _changes(tmp_path, "model.xlsx",
                       [("Commissioning", 1_200_000), ("Financing fees", 2_500_000)],
                       [("Transformers", 1_200_000), ("Commissioning", 2_500_000)])
    assert changes["model.xlsx#Model!B3"].status is ChangeStatus.AMBIGUOUS
    assert not changes["model.xlsx#Model!B4"].mapped


def test_a_renamed_row_taking_over_a_deleted_rows_label_is_not_followed(tmp_path):
    """Seeds G06:0:277 and G06:0:160: the check then reported a false FAIL."""
    changes = _changes(tmp_path, "model.xlsx",
                       [("Owner costs", 3_100_000), ("Water supply", 4_000_000)],
                       [("Owner costs", 4_000_000)])
    assert changes["model.xlsx#Model!B3"].status is ChangeStatus.AMBIGUOUS


def test_two_rows_swapping_labels_are_not_followed_by_label(tmp_path):
    """Seed G06:1:201, on unseen seeds: each link jumped to the other row, two false FAILs."""
    changes = _changes(tmp_path, "model.xlsx",
                       [("Grid connection", 2_500_000), ("Training", 4_000_000)],
                       [("Training", 2_500_000), ("Grid connection", 4_000_000)])
    assert changes["model.xlsx#Model!B3"].status is ChangeStatus.AMBIGUOUS
    assert changes["model.xlsx#Model!B4"].status is ChangeStatus.AMBIGUOUS


def test_a_row_that_simply_moved_is_still_followed(tmp_path):
    changes = _changes(tmp_path, "model.xlsx",
                       [("Civil works", 6_200_000), ("Equipment", 3_100_000)],
                       [("Equipment", 3_100_000), ("Civil works", 6_200_000)])
    civil = changes["model.xlsx#Model!B3"]
    assert civil.mapped and civil.new.node_id == "model.xlsx#Model!B4"


# --- how far a change reaches ----------------------------------------------------

def _node(node_id: str) -> Node:
    return Node(node_id=node_id, source_version="v2", kind=NodeKind.SHEET_CELL,
                selector=LocationSelector(document_id="d", kind=NodeKind.SHEET_CELL),
                evidence_text="synthetic")


def _relation(relation_id: str, rtype: RelationType, method: ExtractionMethod,
              *endpoints: tuple[str, str]) -> Relation:
    return Relation(relation_id=relation_id, type=rtype,
                    endpoints=tuple(Endpoint(role=r, node_id=n) for r, n in endpoints),
                    extraction_method=method, review_state=ReviewState.ACCEPTED,
                    valid_source_versions=("v1",))


def test_a_row_added_to_a_formula_withdraws_reviewed_totals_not_links_to_single_components():
    """Found by the G06 review-cost breakdown: one inserted row withdrew every reviewed link."""
    unchanged = ["narrative-item", "narrative-total", "c1", "c2", "total-cell"]
    changes = [NodeChange(ChangeStatus.UNCHANGED, "same", _node(i), _node(i)) for i in unchanged]
    changes.append(NodeChange(ChangeStatus.ADDED, "new in this version", None, _node("c3")))
    link = _relation("rev:item", RelationType.SAME_QUANTITY_AS, ExtractionMethod.HUMAN_AUTHORED,
                     ("narrative", "narrative-item"), ("workbook", "c1"))
    total = _relation("rev:total", RelationType.SUM_OF, ExtractionMethod.HUMAN_AUTHORED,
                      ("total", "narrative-total"), ("operand", "c1"), ("operand", "c2"))
    grown = _relation("rel:formula", RelationType.SUM_OF, ExtractionMethod.EXPLICIT_FORMULA,
                      ("total", "total-cell"), ("operand", "c1"), ("operand", "c2"),
                      ("operand", "c3"))
    closure = compute_closure(changes, [link, total], [grown],
                              {i: "v2" for i in unchanged + ["c3"]})
    assert {o.old.relation_id: o.disposition for o in closure.outcomes} == {
        "rev:item": Disposition.KEEP, "rev:total": Disposition.REREVIEW}
